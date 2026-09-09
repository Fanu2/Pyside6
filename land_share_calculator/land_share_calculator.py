import sys
from fractions import Fraction
from PySide6.QtCore import Qt
from PySide6.QtGui import QAction
from PySide6.QtWidgets import (QApplication,QMainWindow,QWidget,QVBoxLayout,QHBoxLayout,QGridLayout,QLabel,
QLineEdit,QPushButton,QTableWidget,QTableWidgetItem,QHeaderView,QAbstractItemView,QGroupBox,QSplitter,
QFrame,QMessageBox,QFileDialog,QStatusBar)

COLUMNS=["Owner Name","Share (fraction)","Calculated Area","Remarks"]

def fmt_fraction(x):
    return str(x.numerator) if x.denominator==1 else f"{x.numerator}/{x.denominator}"

def area_string(marlas):
    killa=marlas//160; rem=marlas-killa*160
    kanal=rem//20; rem-=kanal*20
    marla=rem.numerator//rem.denominator
    sarshai=round(float(rem-marla)*9)
    if sarshai>=9: sarshai=0; marla+=1
    if marla>=20: marla=0; kanal+=1
    if kanal>=8: kanal=0; killa+=1
    return f"{killa}K-{kanal}K-{marla}M"+(f"-{sarshai}S" if sarshai else "")

def parse_area(k,m):
    k=int(k or 0); m=int(m or 0)
    if k<0 or m<0 or m>=20: raise ValueError("Kanals must be >= 0 and Marlas must be 0–19.")
    return k*20+m

class Window(QMainWindow):
    def __init__(self):
        super().__init__(); self.setWindowTitle("Punjab Land Share Calculator")
        self.resize(1250,800); self.setMinimumSize(1000,650); self.build_ui(); self.add_owner()

    def build_ui(self):
        root=QVBoxLayout(); root.setContentsMargins(24,20,24,18); root.setSpacing(14)
        central=QWidget(); central.setLayout(root); self.setCentralWidget(central)

        header=QFrame(); header.setObjectName("Header"); hl=QHBoxLayout(header)
        b=QVBoxLayout(); t=QLabel("Punjab Land Share Calculator"); t.setObjectName("Title")
        s=QLabel("Khewat-wise ownership • exact fractional shares • revenue-style area display"); s.setObjectName("Subtitle")
        b.addWidget(t); b.addWidget(s); hl.addLayout(b); hl.addStretch()
        self.badge=QLabel("READY"); self.badge.setObjectName("Badge"); hl.addWidget(self.badge); root.addWidget(header)

        g=QGroupBox("Khewat / Land Details"); gl=QGridLayout(g); gl.setContentsMargins(18,18,18,18)
        self.khewat=QLineEdit(); self.khewat.setPlaceholderText("e.g. 588")
        self.village=QLineEdit(); self.village.setPlaceholderText("Optional")
        self.kanal=QLineEdit(); self.kanal.setPlaceholderText("0")
        self.marla=QLineEdit(); self.marla.setPlaceholderText("0")
        for lab,ed,r,c in [("Khewat No.",self.khewat,0,0),("Village",self.village,0,2),("Total Kanals",self.kanal,1,0),("Total Marlas",self.marla,1,2)]:
            gl.addWidget(QLabel(lab),r,c); gl.addWidget(ed,r,c+1)
        root.addWidget(g)

        split=QSplitter(Qt.Horizontal)
        left=QFrame(); ll=QVBoxLayout(left); ll.setContentsMargins(0,0,0,0)
        bar=QHBoxLayout(); bar.addWidget(QLabel("Owners / Co-sharers")); bar.addStretch()
        add=QPushButton("+ Add Owner"); rem=QPushButton("Remove Selected"); calc=QPushButton("Calculate Shares"); calc.setObjectName("Primary")
        bar.addWidget(add); bar.addWidget(rem); bar.addWidget(calc); ll.addLayout(bar)

        self.table=QTableWidget(0,4); self.table.setHorizontalHeaderLabels(COLUMNS)
        self.table.setSelectionBehavior(QAbstractItemView.SelectRows); self.table.setSelectionMode(QAbstractItemView.ExtendedSelection)
        self.table.setAlternatingRowColors(True); self.table.verticalHeader().setVisible(False)
        h=self.table.horizontalHeader(); h.setSectionResizeMode(0,QHeaderView.Stretch); h.setSectionResizeMode(1,QHeaderView.ResizeToContents)
        h.setSectionResizeMode(2,QHeaderView.ResizeToContents); h.setSectionResizeMode(3,QHeaderView.Stretch); ll.addWidget(self.table)

        cards=QHBoxLayout(); self.cards={}
        for key in ["TOTAL LAND","SHARE TOTAL","REMAINING","OWNERS"]:
            f=QFrame(); f.setObjectName("Card"); q=QVBoxLayout(f); a=QLabel(key); a.setObjectName("CardLabel"); v=QLabel("0"); v.setObjectName("CardValue")
            q.addWidget(a); q.addWidget(v); self.cards[key]=v; cards.addWidget(f)
        ll.addLayout(cards); split.addWidget(left)

        right=QFrame(); rl=QVBoxLayout(right); rl.setContentsMargins(10,0,0,0)
        rg=QGroupBox("Calculation Preview"); rq=QVBoxLayout(rg)
        self.preview=QTableWidget(0,3); self.preview.setHorizontalHeaderLabels(["Owner","Share","Area"]); self.preview.verticalHeader().setVisible(False)
        self.preview.setEditTriggers(QAbstractItemView.NoEditTriggers)
        ph=self.preview.horizontalHeader(); ph.setSectionResizeMode(0,QHeaderView.Stretch); ph.setSectionResizeMode(1,QHeaderView.ResizeToContents); ph.setSectionResizeMode(2,QHeaderView.ResizeToContents)
        rq.addWidget(self.preview); rl.addWidget(rg,1)
        ref=QGroupBox("Quick Reference"); rx=QVBoxLayout(ref)
        lab=QLabel("<b>Area system</b><br>1 Kanal = 20 Marlas<br>1 Killa = 8 Kanals = 160 Marlas<br>1 Marla = 9 Sarshai<br><br><b>Share examples</b><br>1/2 • 1/4 • 3/8 • 5/8")
        lab.setWordWrap(True); rx.addWidget(lab); rl.addWidget(ref)
        split.addWidget(right); split.setSizes([820,380]); root.addWidget(split,1)

        bottom=QHBoxLayout(); sample=QPushButton("Load Example"); clear=QPushButton("Clear All"); export=QPushButton("Export Excel / CSV")
        bottom.addWidget(sample); bottom.addWidget(clear); bottom.addStretch(); bottom.addWidget(export); root.addLayout(bottom)

        add.clicked.connect(self.add_owner); rem.clicked.connect(self.remove); calc.clicked.connect(self.calculate)
        sample.clicked.connect(self.example); clear.clicked.connect(self.clear); export.clicked.connect(self.export)
        self.setStatusBar(QStatusBar()); self.statusBar().showMessage("Ready")
        menu=self.menuBar().addMenu("&File")
        a=QAction("New Calculation",self); a.setShortcut("Ctrl+N"); a.triggered.connect(self.clear); menu.addAction(a)
        a=QAction("Calculate",self); a.setShortcut("F5"); a.triggered.connect(self.calculate); menu.addAction(a)

    def add_owner(self):
        r=self.table.rowCount(); self.table.insertRow(r)
        for c in range(4): self.table.setItem(r,c,QTableWidgetItem(""))
        self.table.setCurrentCell(r,0)

    def remove(self):
        rows=sorted({i.row() for i in self.table.selectedIndexes()},reverse=True)
        if not rows: QMessageBox.information(self,"Remove","Select one or more rows first."); return
        for r in rows: self.table.removeRow(r)
        self.calculate()

    def rows(self):
        return [[self.table.item(r,c).text().strip() if self.table.item(r,c) else "" for c in range(4)] for r in range(self.table.rowCount())]

    def calculate(self):
        try: total=Fraction(parse_area(self.kanal.text(),self.marla.text()))
        except ValueError as e: QMessageBox.warning(self,"Invalid Area",str(e)); return
        parsed=[]; total_share=Fraction(0)
        for name,share_text,_,remarks in self.rows():
            if not name and not share_text: continue
            if not name or not share_text: QMessageBox.warning(self,"Incomplete Owner","Each active row needs a name and share."); return
            try: share=Fraction(share_text)
            except Exception: QMessageBox.warning(self,"Invalid Share",f"Invalid fraction for {name}: {share_text}"); return
            if share<=0: QMessageBox.warning(self,"Invalid Share",f"Share for {name} must be greater than zero."); return
            total_share+=share; parsed.append((name,share,total*share,remarks))
        state="VALID" if total_share==1 else ("OVER" if total_share>1 else "PARTIAL")
        self.badge.setText(state); self.badge.setProperty("state",state); self.badge.style().unpolish(self.badge); self.badge.style().polish(self.badge)
        self.preview.setRowCount(0)
        for name,share,area,_ in parsed:
            r=self.preview.rowCount(); self.preview.insertRow(r)
            self.preview.setItem(r,0,QTableWidgetItem(name)); self.preview.setItem(r,1,QTableWidgetItem(fmt_fraction(share))); self.preview.setItem(r,2,QTableWidgetItem(area_string(area)))
        self.cards["TOTAL LAND"].setText(f"{total//20}-{total%20}"); self.cards["SHARE TOTAL"].setText(fmt_fraction(total_share))
        self.cards["REMAINING"].setText(fmt_fraction(max(Fraction(0),1-total_share))); self.cards["OWNERS"].setText(str(len(parsed)))
        self.table.blockSignals(True)
        for r,(_,st,_,_) in enumerate(self.rows()):
            try: self.table.setItem(r,2,QTableWidgetItem(area_string(total*Fraction(st))) if st else QTableWidgetItem(""))
            except Exception: self.table.setItem(r,2,QTableWidgetItem(""))
        self.table.blockSignals(False)
        self.statusBar().showMessage({"VALID":"Shares total exactly 100%.","OVER":"Shares exceed 100%.","PARTIAL":"Shares are below 100%; remaining share is shown."}[state])

    def clear(self):
        for e in (self.khewat,self.village,self.kanal,self.marla): e.clear()
        self.table.setRowCount(0); self.preview.setRowCount(0); self.badge.setText("READY")
        for k,v in [("TOTAL LAND","0-0"),("SHARE TOTAL","0"),("REMAINING","0"),("OWNERS","0")]: self.cards[k].setText(v)
        self.add_owner()

    def example(self):
        self.clear(); self.khewat.setText("588"); self.village.setText("Taruaana"); self.kanal.setText("8"); self.marla.setText("0")
        data=[("Owner 1","1/2","","Half share"),("Owner 2","1/4","","Quarter share"),("Owner 3","1/4","","Quarter share")]
        self.table.setRowCount(0)
        for row in data:
            r=self.table.rowCount(); self.table.insertRow(r)
            for c,v in enumerate(row): self.table.setItem(r,c,QTableWidgetItem(v))
        self.calculate()

    def export(self):
        path,_=QFileDialog.getSaveFileName(self,"Export Calculation","","Excel Workbook (*.xlsx);;CSV File (*.csv)")
        if not path: return
        try:
            rows=self.rows()
            if path.lower().endswith(".csv"):
                import csv
                with open(path,"w",newline="",encoding="utf-8-sig") as f:
                    w=csv.writer(f); w.writerow(["Khewat No.",self.khewat.text()]); w.writerow(["Village",self.village.text()])
                    w.writerow(["Total Kanals",self.kanal.text()]); w.writerow(["Total Marlas",self.marla.text()]); w.writerow([]); w.writerow(COLUMNS); w.writerows(rows)
            else:
                from openpyxl import Workbook
                from openpyxl.styles import Font,Alignment
                from openpyxl.worksheet.table import Table,TableStyleInfo
                wb=Workbook(); ws=wb.active; ws.title="Land Share Calculation"
                for row in [["Khewat No.",self.khewat.text()],["Village",self.village.text()],["Total Kanals",self.kanal.text()],["Total Marlas",self.marla.text()],[],COLUMNS]+rows: ws.append(row)
                for cell in ws[6]: cell.font=Font(bold=True)
                ws.freeze_panes="A7"
                for col,width in zip("ABCD",[28,18,28,35]): ws.column_dimensions[col].width=width
                for row in ws.iter_rows():
                    for cell in row: cell.alignment=Alignment(vertical="top",wrap_text=True)
                end=6+len(rows)
                if end>=7:
                    tab=Table(displayName="OwnerShares",ref=f"A6:D{end}"); tab.tableStyleInfo=TableStyleInfo(name="TableStyleMedium2",showRowStripes=True); ws.add_table(tab)
                wb.save(path)
            QMessageBox.information(self,"Export Complete",f"Saved successfully:\n{path}")
        except Exception as e: QMessageBox.critical(self,"Export Error",str(e))

STYLE=r"""
QMainWindow,QWidget{background:#f4f7fb;color:#172033}
QMenuBar{background:white;padding:5px}
QFrame#Header{background:#172033;border-radius:16px}
QLabel#Title{color:white;font-size:25px;font-weight:800}
QLabel#Subtitle{color:#cbd5e1;font-size:12px}
QLabel#Badge{background:#e7f7ee;color:#176b3a;border-radius:14px;padding:8px 14px;font-weight:800}
QLabel#Badge[state="OVER"]{background:#fde8e8;color:#a61b1b}
QLabel#Badge[state="PARTIAL"]{background:#fff3cd;color:#8a5a00}
QGroupBox{background:white;border:1px solid #dbe2ec;border-radius:12px;margin-top:10px;font-weight:700;padding-top:12px}
QGroupBox::title{subcontrol-origin:margin;left:16px;padding:0 6px}
QLineEdit{background:#fbfcfe;border:1px solid #cfd8e3;border-radius:7px;padding:8px 10px}
QPushButton{background:white;border:1px solid #cfd8e3;border-radius:8px;padding:8px 14px;font-weight:650}
QPushButton:hover{background:#edf3fb}
QPushButton#Primary{background:#284b7a;color:white;border:0}
QTableWidget{background:white;border:1px solid #dbe2ec;border-radius:9px;gridline-color:#e8edf3;selection-background-color:#dce8f7;selection-color:#172033}
QHeaderView::section{background:#eef2f7;border:0;border-bottom:1px solid #dbe2ec;padding:9px;font-weight:750}
QFrame#Card{background:white;border:1px solid #dbe2ec;border-radius:10px}
QLabel#CardLabel{color:#697586;font-size:10px;font-weight:750}
QLabel#CardValue{color:#172033;font-size:18px;font-weight:800}
QStatusBar{background:white}
"""

if __name__=="__main__":
    app=QApplication(sys.argv); app.setStyleSheet(STYLE); app.setApplicationName("Punjab Land Share Calculator")
    win=Window(); win.show(); sys.exit(app.exec())

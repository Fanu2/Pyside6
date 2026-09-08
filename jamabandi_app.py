import sys
import sqlite3
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo
from PySide6.QtCore import Qt
from PySide6.QtGui import QAction
from PySide6.QtWidgets import (
    QApplication, QMainWindow, QWidget, QVBoxLayout, QHBoxLayout,
    QFormLayout, QLineEdit, QPushButton, QTableWidget, QTableWidgetItem,
    QFileDialog, QMessageBox, QLabel, QDialog, QHeaderView, QAbstractItemView,
    QGroupBox, QSplitter, QToolBar
)

COLUMNS = [
    "Khewat / Jamabandi No.",
    "Khatauni No.",
    "Name / Patti",
    "Owner Name",
    "Cultivator",
    "Irrigation Source",
    "Khasra / Murabba / Killa No.",
    "Area / Land Type",
]

APP_DIR = Path(__file__).resolve().parent
SAMPLE_XLSX = APP_DIR / "jamabandi_sample_all_first_8.xlsx"
DB_FILE = APP_DIR / "jamabandi.db"


class RecordDialog(QDialog):
    def __init__(self, parent=None, record=None):
        super().__init__(parent)
        self.setWindowTitle("New Record" if record is None else "Edit Record")
        self.resize(720, 500)
        layout = QFormLayout(self)
        self.edits = []
        for i, col in enumerate(COLUMNS):
            edit = QLineEdit(record[i] if record else "")
            edit.setClearButtonEnabled(True)
            layout.addRow(col + ":", edit)
            self.edits.append(edit)

        buttons = QHBoxLayout()
        save = QPushButton("Save")
        cancel = QPushButton("Cancel")
        save.clicked.connect(self.accept)
        cancel.clicked.connect(self.reject)
        buttons.addWidget(save)
        buttons.addWidget(cancel)
        layout.addRow(buttons)

    def record(self):
        return [e.text().strip() for e in self.edits]


class JamabandiDB:
    def __init__(self, path):
        self.path = path
        self.conn = sqlite3.connect(path)
        self.conn.execute("""
            CREATE TABLE IF NOT EXISTS records (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                c1 TEXT, c2 TEXT, c3 TEXT, c4 TEXT,
                c5 TEXT, c6 TEXT, c7 TEXT, c8 TEXT
            )
        """)
        self.conn.commit()

    def all(self):
        return self.conn.execute(
            "SELECT id,c1,c2,c3,c4,c5,c6,c7,c8 FROM records ORDER BY id"
        ).fetchall()

    def insert(self, r):
        self.conn.execute(
            "INSERT INTO records(c1,c2,c3,c4,c5,c6,c7,c8) VALUES(?,?,?,?,?,?,?,?)",
            r
        )
        self.conn.commit()

    def update(self, rid, r):
        self.conn.execute(
            "UPDATE records SET c1=?,c2=?,c3=?,c4=?,c5=?,c6=?,c7=?,c8=? WHERE id=?",
            (*r, rid)
        )
        self.conn.commit()

    def delete(self, rid):
        self.conn.execute("DELETE FROM records WHERE id=?", (rid,))
        self.conn.commit()

    def replace_all(self, records):
        with self.conn:
            self.conn.execute("DELETE FROM records")
            self.conn.executemany(
                "INSERT INTO records(c1,c2,c3,c4,c5,c6,c7,c8) VALUES(?,?,?,?,?,?,?,?)",
                records
            )


class MainWindow(QMainWindow):
    def __init__(self):
        super().__init__()
        self.setWindowTitle("Jamabandi Record Manager — First 8 Columns")
        self.resize(1550, 850)
        self.db = JamabandiDB(DB_FILE)
        self.records = []
        self.visible_ids = []
        self.build_ui()

        # First run: load the complete supplied sample dataset.
        if not self.db.all() and SAMPLE_XLSX.exists():
            self.import_from_path(SAMPLE_XLSX)
        self.refresh()

    def build_ui(self):
        central = QWidget()
        self.setCentralWidget(central)
        root = QVBoxLayout(central)

        title = QLabel("Jamabandi Record Manager")
        title.setStyleSheet("font-size:24px;font-weight:700;padding:4px;")
        root.addWidget(title)

        meta = QLabel(
            "तारूआना  |  हदबस्त 315  |  कालांवाली  |  सिरसा  |  2023-2024"
        )
        meta.setStyleSheet("color:#555;padding-bottom:6px;")
        root.addWidget(meta)

        toolbar = QToolBar()
        toolbar.setMovable(False)
        root.addWidget(toolbar)

        actions = [
            ("New", self.new_record, "Ctrl+N"),
            ("Edit", self.edit_selected, "Ctrl+E"),
            ("Delete", self.delete_selected, "Delete"),
            ("Import XLSX", self.import_xlsx, "Ctrl+O"),
            ("Export XLSX", self.export_xlsx, "Ctrl+S"),
            ("Reload", self.refresh, "F5"),
        ]
        for text, slot, shortcut in actions:
            action = QAction(text, self)
            action.setShortcut(shortcut)
            action.triggered.connect(slot)
            toolbar.addAction(action)

        search_group = QGroupBox("Search")
        sh = QHBoxLayout(search_group)
        self.search = QLineEdit()
        self.search.setPlaceholderText(
            "Search Khewat, Khatauni, Patti, Owner, Cultivator, Irrigation, Khasra or Area…"
        )
        self.search.textChanged.connect(self.refresh)
        sh.addWidget(self.search)
        clear = QPushButton("Clear")
        clear.clicked.connect(self.search.clear)
        sh.addWidget(clear)
        root.addWidget(search_group)

        self.table = QTableWidget(0, 8)
        self.table.setHorizontalHeaderLabels(COLUMNS)
        self.table.setSelectionBehavior(QAbstractItemView.SelectRows)
        self.table.setSelectionMode(QAbstractItemView.SingleSelection)
        self.table.setEditTriggers(QAbstractItemView.NoEditTriggers)
        self.table.setAlternatingRowColors(True)
        self.table.setSortingEnabled(True)
        self.table.doubleClicked.connect(lambda _: self.edit_selected())
        self.table.horizontalHeader().setSectionResizeMode(QHeaderView.Interactive)
        self.table.horizontalHeader().setStretchLastSection(True)
        self.table.verticalHeader().setDefaultSectionSize(32)
        root.addWidget(self.table, 1)

        self.status = QLabel()
        self.status.setStyleSheet("padding:5px;")
        root.addWidget(self.status)

        file_menu = self.menuBar().addMenu("&File")
        for text, slot, shortcut in [
            ("New", self.new_record, "Ctrl+N"),
            ("Import XLSX", self.import_xlsx, "Ctrl+O"),
            ("Export XLSX", self.export_xlsx, "Ctrl+S"),
            ("Exit", self.close, "Ctrl+Q"),
        ]:
            a = QAction(text, self)
            a.setShortcut(shortcut)
            a.triggered.connect(slot)
            file_menu.addAction(a)

    def refresh(self):
        all_rows = self.db.all()
        q = self.search.text().strip().casefold()
        if q:
            rows = [
                row for row in all_rows
                if q in " | ".join(str(x or "") for x in row[1:]).casefold()
            ]
        else:
            rows = all_rows

        self.records = {row[0]: list(row[1:]) for row in rows}
        self.visible_ids = [row[0] for row in rows]

        self.table.setSortingEnabled(False)
        self.table.setRowCount(len(rows))
        for vr, row in enumerate(rows):
            for c, value in enumerate(row[1:]):
                item = QTableWidgetItem(value or "")
                item.setToolTip(value or "")
                self.table.setItem(vr, c, item)
        self.table.setSortingEnabled(True)

        self.status.setText(
            f"Total records: {len(all_rows)}    |    Displayed: {len(rows)}"
        )

    def selected_id(self):
        row = self.table.currentRow()
        if row < 0 or row >= len(self.visible_ids):
            return None
        return self.visible_ids[row]

    def new_record(self):
        dlg = RecordDialog(self)
        if dlg.exec() == QDialog.Accepted:
            self.db.insert(dlg.record())
            self.refresh()

    def edit_selected(self):
        rid = self.selected_id()
        if rid is None:
            QMessageBox.information(self, "Select record", "Select a record first.")
            return
        dlg = RecordDialog(self, self.records[rid])
        if dlg.exec() == QDialog.Accepted:
            self.db.update(rid, dlg.record())
            self.refresh()

    def delete_selected(self):
        rid = self.selected_id()
        if rid is None:
            QMessageBox.information(self, "Select record", "Select a record first.")
            return
        if QMessageBox.question(
            self, "Delete record",
            "Delete the selected record?"
        ) == QMessageBox.Yes:
            self.db.delete(rid)
            self.refresh()

    def import_xlsx(self):
        path, _ = QFileDialog.getOpenFileName(
            self, "Import Excel Workbook", "", "Excel Workbook (*.xlsx)"
        )
        if path:
            self.import_from_path(Path(path))

    def import_from_path(self, path):
        try:
            wb = load_workbook(path, read_only=True, data_only=True)
            ws = wb[wb.sheetnames[0]]
            rows = list(ws.iter_rows(values_only=True))
            if not rows:
                raise ValueError("The workbook is empty.")

            headers = [str(x).strip() if x is not None else "" for x in rows[0]]
            # Match by header names when possible; otherwise use first 8 columns.
            indexes = []
            for col in COLUMNS:
                if col in headers:
                    indexes.append(headers.index(col))
                else:
                    indexes = list(range(min(8, len(headers))))
                    break

            records = []
            for row in rows[1:]:
                rec = []
                for idx in indexes[:8]:
                    rec.append("" if idx >= len(row) or row[idx] is None else str(row[idx]))
                rec += [""] * (8 - len(rec))
                if any(x.strip() for x in rec):
                    records.append(rec)

            wb.close()
            self.db.replace_all(records)
            self.search.clear()
            self.refresh()
            QMessageBox.information(
                self, "Import complete",
                f"Imported {len(records)} records using the first 8 application columns."
            )
        except Exception as e:
            QMessageBox.critical(self, "Import failed", str(e))

    def export_xlsx(self):
        path, _ = QFileDialog.getSaveFileName(
            self, "Export Excel Workbook",
            "jamabandi_export.xlsx",
            "Excel Workbook (*.xlsx)"
        )
        if not path:
            return

        try:
            wb = Workbook()
            ws = wb.active
            ws.title = "Jamabandi Data"
            ws.append(COLUMNS)

            # Export the complete database, not just the current filtered view.
            for row in self.db.all():
                ws.append(list(row[1:]))

            fill = PatternFill("solid", fgColor="1F4E78")
            for c in ws[1]:
                c.fill = fill
                c.font = Font(color="FFFFFF", bold=True)
                c.alignment = Alignment(
                    horizontal="center", vertical="center", wrap_text=True
                )

            widths = [22, 18, 22, 34, 34, 28, 32, 32]
            for i, width in enumerate(widths, 1):
                ws.column_dimensions[chr(64+i)].width = width
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions

            if ws.max_row >= 2:
                table = Table(displayName="JamabandiExport", ref=ws.dimensions)
                table.tableStyleInfo = TableStyleInfo(
                    name="TableStyleMedium2", showRowStripes=True
                )
                ws.add_table(table)

            wb.save(path)
            QMessageBox.information(
                self, "Export complete",
                f"Exported {ws.max_row - 1} records to:\n{path}"
            )
        except Exception as e:
            QMessageBox.critical(self, "Export failed", str(e))


if __name__ == "__main__":
    app = QApplication(sys.argv)
    app.setStyle("Fusion")
    win = MainWindow()
    win.show()
    sys.exit(app.exec())
